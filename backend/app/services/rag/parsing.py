from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass

from app.services.security.injection import sanitize, strip_html


@dataclass
class Section:
    heading: str
    text: str
    page: int | None = None


class UnsupportedDocument(ValueError):
    pass


def detect_kind(filename: str, mime: str) -> str:
    name = filename.lower()
    if name.endswith(".pdf") or mime == "application/pdf":
        return "pdf"
    if name.endswith(".docx") or "wordprocessingml" in mime:
        return "docx"
    if name.endswith((".html", ".htm")) or mime == "text/html":
        return "html"
    if name.endswith(".csv") or mime == "text/csv":
        return "csv"
    if name.endswith((".md", ".markdown")):
        return "markdown"
    if name.endswith((".txt", ".text", ".json")) or mime.startswith("text/"):
        return "text"
    raise UnsupportedDocument(f"unsupported document type: {filename} ({mime})")


def _markdown_sections(text: str) -> list[Section]:
    sections: list[Section] = []
    heading, buf = "", []
    for line in text.splitlines():
        m = re.match(r"^\s{0,3}(#{1,6})\s+(.*)$", line)
        if m:
            if "".join(buf).strip():
                sections.append(Section(heading, "\n".join(buf).strip()))
            heading, buf = m.group(2).strip(), []
        else:
            buf.append(line)
    if "".join(buf).strip():
        sections.append(Section(heading, "\n".join(buf).strip()))
    return sections


def parse_document(data: bytes, filename: str, mime: str = "") -> list[Section]:
    kind = detect_kind(filename, mime)
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return [Section(f"Page {i + 1}", sanitize(p.extract_text() or ""), i + 1)
                for i, p in enumerate(reader.pages) if (p.extract_text() or "").strip()]
    if kind == "docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        sections: list[Section] = []
        heading, buf = "", []
        for para in d.paragraphs:
            if para.style is not None and para.style.name.lower().startswith("heading"):
                if buf:
                    sections.append(Section(heading, sanitize("\n".join(buf))))
                heading, buf = para.text.strip(), []
            elif para.text.strip():
                buf.append(para.text)
        for table in d.tables:
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
            buf.append("\n".join(rows))
        if buf:
            sections.append(Section(heading, sanitize("\n".join(buf))))
        return sections
    text = data.decode("utf-8", errors="replace")
    if kind == "html":
        return _markdown_sections(strip_html(text))
    if kind == "csv":
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            return []
        header = rows[0]
        lines = ["; ".join(f"{h}: {v}" for h, v in zip(header, r, strict=False)) for r in rows[1:]]
        return [Section("Rows", sanitize("\n".join(lines)))]
    return _markdown_sections(sanitize(text))
